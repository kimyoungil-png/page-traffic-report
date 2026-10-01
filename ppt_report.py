from __future__ import annotations

import io
import os
import tempfile
from copy import deepcopy
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt

R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


ROOT = Path(__file__).resolve().parent
DEFAULT_TEMPLATE_PATH = ROOT / "templates" / "explore_dotcom_sample_v2.pptx"
FONT_FACE = "Meiryo UI"

BLUE = RGBColor(30, 53, 227)
RATIO_BLUE = RGBColor(0, 49, 244)
RED = RGBColor(255, 0, 0)
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(128, 128, 128)
PALE = RGBColor(181, 181, 181)
LIGHT_PLACEHOLDER = RGBColor(246, 246, 246)
PLACEHOLDER_LINE = RGBColor(210, 210, 210)


def _load_template() -> Presentation:
    path = Path(os.getenv("PPT_TEMPLATE_PATH", str(DEFAULT_TEMPLATE_PATH)))
    if not path.exists():
        raise RuntimeError(f"PowerPointテンプレートが見つかりません: {path}")
    return Presentation(str(path))


def _remap_relationship_ids(source_slide, new_slide, cloned_element) -> None:
    """Copy relationship targets used inside a cloned shape."""
    rid_map: dict[str, str] = {}
    for element in cloned_element.iter():
        for attr_name in (R_NS + "embed", R_NS + "link", R_NS + "id"):
            old_rid = element.get(attr_name)
            if not old_rid or old_rid not in source_slide.part.rels:
                continue
            if old_rid not in rid_map:
                rel = source_slide.part.rels[old_rid]
                rid_map[old_rid] = new_slide.part.relate_to(rel._target, rel.reltype)
            element.set(attr_name, rid_map[old_rid])


def _duplicate_template_slide(presentation: Presentation, source_slide):
    new_slide = presentation.slides.add_slide(source_slide.slide_layout)
    for shape in list(new_slide.shapes):
        element = shape.element
        element.getparent().remove(element)
    for shape in source_slide.shapes:
        cloned = deepcopy(shape.element)
        _remap_relationship_ids(source_slide, new_slide, cloned)
        new_slide.shapes._spTree.insert_element_before(cloned, "p:extLst")
    return new_slide


def _remove_shape(shape) -> None:
    element = shape.element
    element.getparent().remove(element)


def _clear_text_frame(shape):
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.NONE
    return text_frame


def _set_font(run, size: float, color=BLACK, bold: bool = False, italic: bool = False):
    font = run.font
    font.name = FONT_FACE
    font.size = Pt(size)
    font.bold = bold
    font.italic = italic
    font.color.rgb = color


def _set_simple_text(shape, text: str, size: float, color=BLACK, bold: bool = False, italic: bool = False, align=None):
    tf = _clear_text_frame(shape)
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text or ""
    _set_font(run, size, color, bold, italic)


def _clear_paragraph_preserve_style(paragraph) -> None:
    p_element = paragraph._p
    for child in list(p_element):
        if child.tag.endswith('}pPr'):
            continue
        p_element.remove(child)


def _ensure_styled_paragraph(text_frame, index: int):
    while len(text_frame.paragraphs) <= index:
        new_p = text_frame.add_paragraph()
        base_p_pr = text_frame.paragraphs[0]._p.pPr
        if base_p_pr is not None and new_p._p.pPr is not None:
            new_p._p.remove(new_p._p.pPr)
            new_p._p.insert(0, deepcopy(base_p_pr))
        elif base_p_pr is not None:
            new_p._p.insert(0, deepcopy(base_p_pr))
    return text_frame.paragraphs[index]


def _set_paragraph_text_preserve_style(paragraph, text: str, size: float, color=BLACK, bold: bool = False):
    _clear_paragraph_preserve_style(paragraph)
    run = paragraph.add_run()
    run.text = str(text or "")
    _set_font(run, size, color, bold)


def _force_no_autofit(slide) -> None:
    for shape in slide.shapes:
        if hasattr(shape, "text_frame") and shape.has_text_frame:
            shape.text_frame.auto_size = MSO_AUTO_SIZE.NONE


def _ratio_color(ratio_label: str):
    if not ratio_label or ratio_label in ("—", "NEW"):
        return GREY
    try:
        value = float(str(ratio_label).replace("x", ""))
    except ValueError:
        return RATIO_BLUE
    return RED if value < 1 else RATIO_BLUE


def _fmt_num(value) -> str:
    if value is None or value == "":
        return "—"
    try:
        return f"{int(round(float(value))):,}"
    except Exception:
        return str(value)


def _fmt_pct(value, digits: int = 0) -> str:
    if value is None or value == "":
        return "—"
    try:
        return f"{float(value):.{digits}f}%"
    except Exception:
        return str(value)


def _find_text_shape(slide, predicate):
    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        text = shape.text or ""
        if predicate(text, shape):
            return shape
    return None


def _find_traffic_table(slide):
    for shape in slide.shapes:
        if shape.has_table and len(shape.table.rows) == 11 and len(shape.table.columns) == 9:
            return shape.table
    return None


def _find_gsc_table(slide):
    for shape in slide.shapes:
        if shape.has_table and len(shape.table.rows) == 12 and len(shape.table.columns) == 3:
            return shape.table
    return None


def _set_cell_text(cell, text: str, size: float = 8.0, color=BLACK, bold: bool = False, align=PP_ALIGN.CENTER):
    cell.text = ""
    tf = cell.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text or "")
    _set_font(run, size, color, bold, False)


def _replace_header(slide, page: dict[str, Any]) -> None:
    breadcrumb = _find_text_shape(slide, lambda text, _s: ">" in text and ("support" in text or "hint" in text or "special" in text))
    if breadcrumb is not None:
        _set_simple_text(breadcrumb, page.get("breadcrumb", ""), 10.0, PALE, True)

    title_shape = _find_text_shape(slide, lambda text, _s: "今週Entry Visit" in text or "今週 Entry Visit" in text)
    if title_shape is not None:
        tf = _clear_text_frame(title_shape)
        p0 = tf.paragraphs[0]
        p0.space_after = Pt(0)
        r0 = p0.add_run()
        r0.text = page.get("report_path", "")
        _set_font(r0, 16.0, BLACK, True)

        p1 = tf.add_paragraph()
        p1.space_before = Pt(0)
        p1.space_after = Pt(0)
        r1 = p1.add_run()
        r1.text = f"今週Entry Visit {page.get('total_current_compact', '0')} ("
        _set_font(r1, 15.0, BLACK, True)

        ratio = page.get("total_ratio_label", "—")
        r2 = p1.add_run()
        r2.text = str(ratio)
        _set_font(r2, 14.0, _ratio_color(str(ratio)), True)

        r3 = p1.add_run()
        r3.text = " vs 先週) "
        _set_font(r3, 14.0, BLACK, True)

        r4 = p1.add_run()
        r4.text = f"→{page.get('headline_comment', '')}"
        _set_font(r4, 14.0, BLUE, True)

    bullet_shape = _find_text_shape(slide, lambda text, _s: "流入割合" in text)
    if bullet_shape is not None:
        tf = bullet_shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE

        p = _ensure_styled_paragraph(tf, 0)
        _set_paragraph_text_preserve_style(p, page.get("share_line", ""), 11.0, BLACK)

        secondary = page.get("secondary_line", "") or page.get("detail_comment", "")
        p2 = _ensure_styled_paragraph(tf, 1)
        _set_paragraph_text_preserve_style(p2, secondary, 11.0, BLACK)

        for extra in tf.paragraphs[2:]:
            _set_paragraph_text_preserve_style(extra, "", 11.0, BLACK)

    meta_shape = _find_text_shape(slide, lambda text, _s: text.startswith("Title :"))
    if meta_shape is not None:
        meta_title = page.get("meta_title") or "—"
        _set_simple_text(meta_shape, f"Title : {meta_title}", 10.6, BLACK, True, True)

    footer_shape = _find_text_shape(slide, lambda text, _s: "Data：" in text and "samsung.com" in text)
    if footer_shape is not None:
        footer_text = f"Data：{page.get('period_label', '')}\n{page.get('url', '')}"
        _set_simple_text(footer_shape, footer_text, 6.4, GREY)


def _replace_traffic_table(slide, rows: list[dict[str, Any]]) -> None:
    table = _find_traffic_table(slide)
    if table is None:
        return

    for r, data in enumerate((rows or [])[:9], start=2):
        is_total = str(data.get("channel") or "") == "Total"
        values = [
            data.get("channel", ""),
            _fmt_num(data.get("previous_entry")),
            _fmt_num(data.get("previous_cta")),
            _fmt_pct(data.get("previous_bounce"), 0),
            _fmt_num(data.get("current_entry")),
            data.get("ratio") or "—",
            _fmt_num(data.get("current_cta")),
            _fmt_pct(data.get("current_ctr"), 1),
            _fmt_pct(data.get("current_bounce"), 0),
        ]
        for c, value in enumerate(values):
            color = BLACK
            size = 8.0
            if c in (1, 2, 3):
                color = GREY
            if c == 5:
                color = _ratio_color(str(value))
                size = 7.6
            if c == 0:
                size = 7.4
            _set_cell_text(table.cell(r, c), value, size=size, color=color, bold=is_total)


def _replace_gsc_table(slide, queries: list[dict[str, Any]]) -> None:
    table = _find_gsc_table(slide)
    if table is None:
        return

    for i in range(10):
        row = i + 2
        query = queries[i] if i < len(queries or []) else {}
        q_text = str(query.get("query") or "")
        clicks = ""
        if q_text:
            raw_clicks = query.get("clicks", "")
            try:
                clicks = str(int(round(float(raw_clicks))))
            except Exception:
                clicks = str(raw_clicks or "")
        _set_cell_text(table.cell(row, 0), str(i + 1), size=7.2)
        _set_cell_text(table.cell(row, 1), q_text, size=7.0)
        _set_cell_text(table.cell(row, 2), clicks, size=7.0)


def _replace_screenshot(slide, screenshot_bytes: bytes | None) -> None:
    target = None
    for shape in slide.shapes:
        text = (getattr(shape, "text", "") or "").strip() if hasattr(shape, "text") else ""
        name = str(getattr(shape, "name", "") or "")
        if text == "Mobile Screenshot" or name == "Rectangle 31":
            target = shape
            break

    if target is None:
        for shape in slide.shapes:
            if shape.shape_type == 13 and shape.left > 450000 and shape.top > 2300000 and shape.width > 1500000 and shape.height > 3000000:
                target = shape
                break

    if target is not None:
        left, top, width, height = target.left, target.top, target.width, target.height
        _remove_shape(target)
    else:
        left, top, width, height = 520638, 2452688, 1993900, 4000500

    if screenshot_bytes:
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
            f.write(screenshot_bytes)
            image_path = f.name
        try:
            slide.shapes.add_picture(image_path, left, top, width=width, height=height)
        finally:
            try:
                os.remove(image_path)
            except FileNotFoundError:
                pass
    else:
        rect = slide.shapes.add_shape(1, left, top, width, height)
        rect.fill.solid()
        rect.fill.fore_color.rgb = LIGHT_PLACEHOLDER
        rect.line.color.rgb = PLACEHOLDER_LINE
        tf = rect.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(0)
        run = p.add_run()
        run.text = "Mobile Screenshot"
        _set_font(run, 9.0, RGBColor(170, 170, 170))


def _fill_slide(slide, page: dict[str, Any]) -> None:
    _replace_header(slide, page)
    _replace_traffic_table(slide, page.get("table_rows") or [])
    _replace_gsc_table(slide, page.get("gsc_queries") or [])
    _replace_screenshot(slide, page.get("screenshot_bytes"))
    _force_no_autofit(slide)


def build_ppt_report(pages: list[dict[str, Any]]) -> bytes:
    if not pages:
        raise RuntimeError("PowerPointに出力するページがありません。")

    presentation = _load_template()
    template_slide = presentation.slides[0]

    while len(presentation.slides) < len(pages):
        _duplicate_template_slide(presentation, template_slide)

    for idx, page in enumerate(pages):
        _fill_slide(presentation.slides[idx], page)

    output = io.BytesIO()
    presentation.save(output)
    output.seek(0)
    return output.read()
